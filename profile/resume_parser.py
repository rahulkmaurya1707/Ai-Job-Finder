import io
from pathlib import Path
from typing import Union
import pdfplumber
import docx


def extract_text_from_pdf(source: Union[str, Path, io.BytesIO, bytes]) -> str:
    """Extract raw text from a PDF file path or byte stream using pdfplumber."""
    text_content = []
    if isinstance(source, bytes):
        source = io.BytesIO(source)

    with pdfplumber.open(source) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_content.append(page_text)
    return "\n".join(text_content).strip()


def extract_text_from_docx(source: Union[str, Path, io.BytesIO, bytes]) -> str:
    """Extract raw text from a DOCX file path or byte stream using python-docx."""
    if isinstance(source, bytes):
        source = io.BytesIO(source)

    doc = docx.Document(source)
    text_content = []
    for paragraph in doc.paragraphs:
        if paragraph.text:
            text_content.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_content.append(" | ".join(row_text))
    return "\n".join(text_content).strip()


def extract_text_from_resume(
    source: Union[str, Path, io.BytesIO, bytes], filename_or_ext: str
) -> str:
    """Detect file format by extension (.pdf or .docx) and return extracted text."""
    ext = Path(filename_or_ext).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(source)
    elif ext == ".docx":
        return extract_text_from_docx(source)
    else:
        raise ValueError(f"Unsupported file format: {ext}. Expected .pdf or .docx")
